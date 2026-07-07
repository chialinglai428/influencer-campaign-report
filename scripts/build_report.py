"""
建立網紅專案結案成效總表(xlsx)
用法: 在程式中準備好 data_rows 後呼叫 build(data_rows, output_path)

data_rows 每筆為 dict,鍵值:
    name        str  網紅名稱(必填)
    followers   int  粉絲數(必填)
    reach       int  觸及帳號數(必填,依優先序判斷來源,見 SKILL.md)
    likes       int  愛心數(必填)
    comments    int  留言數(必填)
    shares      int  轉發數 = 轉發(repost) + 傳送(send) 相加後的總和(必填)
    saves       int  收藏數(必填)
    link_clicks int 或 None  連結點擊數(optional,無資料則傳 None,不計入加總)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADERS = ["網紅名稱", "粉絲數", "觸及帳號數", "愛心數", "留言數",
           "轉發數(轉發+傳送)", "收藏數", "互動率", "連結點擊數"]

FIELD_ORDER = ["name", "followers", "reach", "likes", "comments", "shares", "saves", "link_clicks"]


def build(data_rows, output_path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "網紅成效總表"

    sheet.append(HEADERS)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="4B0082")
    for col in range(1, len(HEADERS) + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in data_rows:
        values = [row.get(f) for f in FIELD_ORDER]
        # 插入互動率欄位(第8欄)佔位,稍後用公式覆蓋
        values.insert(7, None)
        sheet.append(values)

    n = len(data_rows)
    last_row = 1 + n

    for r in range(2, last_row + 1):
        sheet.cell(row=r, column=8).value = f"=(D{r}+E{r}+F{r})/C{r}"
        sheet.cell(row=r, column=8).number_format = "0.0%"

    total_row = last_row + 1
    sheet.cell(row=total_row, column=1).value = "總計"
    sheet.cell(row=total_row, column=1).font = Font(name="Arial", bold=True)

    for col, letter in zip([2, 3, 4, 5, 6, 7, 9], ["B", "C", "D", "E", "F", "G", "I"]):
        sheet.cell(row=total_row, column=col).value = f"=SUM({letter}2:{letter}{last_row})"
        sheet.cell(row=total_row, column=col).font = Font(name="Arial", bold=True)

    # 加權平均互動率:以觸及帳號數加權,避免不同網紅觸及量差異造成失真
    sheet.cell(row=total_row, column=8).value = (
        f"=SUMPRODUCT(D2:D{last_row}+E2:E{last_row}+F2:F{last_row})/C{total_row}"
    )
    sheet.cell(row=total_row, column=8).number_format = "0.0%"
    sheet.cell(row=total_row, column=8).font = Font(name="Arial", bold=True)

    fill_total = PatternFill("solid", start_color="E6E0F8")
    for col in range(1, 10):
        sheet.cell(row=total_row, column=col).fill = fill_total
        sheet.column_dimensions[get_column_letter(col)].width = 18

    for row in sheet.iter_rows(min_row=2, max_row=total_row, min_col=1, max_col=9):
        for cell in row:
            cell.alignment = Alignment(horizontal="center") if cell.column > 1 else Alignment(horizontal="left")

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    import sys
    print("此腳本需以模組方式匯入使用,範例:\n"
          "from build_report import build\n"
          "build(data_rows, 'output.xlsx')")
